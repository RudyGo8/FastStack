"""Ingest official SPU and SKU mapping table from docs/SPU与物料编码映射表.xlsx."""

from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app.modules.sop.database import SessionLocal
from app.modules.sop.domain.ingestion import SopIngestionService
from app.modules.sop.models.db_sop import SopSpuMapping
from app.modules.sop.schemas.sop import (
    SopImportMeta,
    SopSpuImportRequest,
    SopSpuImportRow,
)
from app.modules.sop.utils.log import get_logger

logger = get_logger(__name__)


def ingest_official_spu_mapping_excel(db: Session, excel_path: str | Path | None = None) -> dict[str, int]:
    """读取 docs/SPU与物料编码映射表.xlsx 并将 162 款 SPU 及 1000+ SKU 映射导入数据库."""
    if excel_path is None:
        for candidate in [
            Path("docs/SPU与物料编码映射表.xlsx"),
            Path(__file__).resolve().parents[4] / "docs" / "SPU与物料编码映射表.xlsx",
        ]:
            if candidate.exists():
                excel_path = candidate
                break

    if not excel_path or not Path(excel_path).exists():
        logger.warning("official_spu_mapping_excel_not_found", path=str(excel_path))
        return {"spus_imported": 0, "mappings_imported": 0}

    wb = openpyxl.load_workbook(str(excel_path))
    sheet_name = "映射表整理 -筛选BU" if "映射表整理 -筛选BU" in wb.sheetnames else wb.sheetnames[0]
    sheet = wb[sheet_name]

    spu_dict: dict[str, SopSpuImportRow] = {}
    sku_mappings: list[tuple[str, str, str]] = []  # (spu_code, sku_code, sku_name)

    for r in range(2, sheet.max_row + 1):
        spu = sheet.cell(r, 1).value
        if not spu:
            continue
        spu_code = str(spu).strip()
        if not spu_code:
            continue

        pline = str(sheet.cell(r, 2).value or "智能硬件").strip()
        bu = str(sheet.cell(r, 3).value or "标准BU").strip()
        brand = str(sheet.cell(r, 4).value or "Magene").strip()
        pcode = str(sheet.cell(r, 5).value or "").strip()
        mname = str(sheet.cell(r, 6).value or "").strip()

        if spu_code not in spu_dict:
            spu_dict[spu_code] = SopSpuImportRow(
                spu_code=spu_code,
                source_spu_code=pcode or spu_code,
                spu_name=mname or f"{brand} {spu_code}",
                product_line=pline,
                brand=brand,
                category=bu,
                lifecycle_stage="成长期",
                mapping_type="sku_to_spu" if pcode else "direct",
                mapping_version="2026_OFFICIAL_MAP",
            )
        if pcode:
            sku_mappings.append((spu_code, pcode, mname))

    now = datetime.now()
    batch_id = f"excel_map_{now.strftime('%Y%m%d_%H%M%S')}"

    # 1. 导入 SPU 主数据
    ingestion = SopIngestionService(db)
    result = ingestion.import_spus(
        SopSpuImportRequest(
            meta=SopImportMeta(
                source_system="ERP_PLM_EXCEL",
                source_table="SPU与物料编码映射表.xlsx",
                batch_id=batch_id,
                snapshot_at=now,
            ),
            rows=list(spu_dict.values()),
        )
    )

    # 2. 批量写入全量 SKU 映射明细
    mappings_count = 0
    for spu_code, sku_code, _mname in sku_mappings:
        existing = (
            db.query(SopSpuMapping)
            .filter(
                SopSpuMapping.source_system == "ERP_PLM_EXCEL",
                SopSpuMapping.source_spu_code == sku_code,
            )
            .first()
        )
        if existing:
            existing.canonical_spu_code = spu_code
            existing.mapping_type = "sku_to_spu"
            existing.mapping_version = "2026_OFFICIAL_MAP"
        else:
            db.add(
                SopSpuMapping(
                    source_system="ERP_PLM_EXCEL",
                    source_spu_code=sku_code,
                    canonical_spu_code=spu_code,
                    mapping_type="sku_to_spu",
                    mapping_version="2026_OFFICIAL_MAP",
                )
            )
        mappings_count += 1

    db.commit()

    logger.info(
        "official_spu_mapping_excel_ingested",
        spus_accepted=result.accepted,
        total_sku_mappings=mappings_count,
    )
    return {
        "spus_imported": result.accepted,
        "mappings_imported": mappings_count,
    }


if __name__ == "__main__":
    db = SessionLocal()
    try:
        res = ingest_official_spu_mapping_excel(db)
        print("Import result:", res)
    finally:
        db.close()
