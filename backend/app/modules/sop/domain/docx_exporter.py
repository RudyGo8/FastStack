"""Generate a fact-only S&OP report from the public report contract."""

import io
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def get_report_readiness_label(completeness_status: str) -> str:
    return {
        "complete": "核心数据已就绪（可评审）",
        "partial": "数据域不完整（待补充）",
        "not_ready": "核心数据未就绪（暂不可评审）",
    }.get(completeness_status, "核心数据未就绪（暂不可评审）")


def _set_cell_background(cell, fill_hex: str) -> None:
    cell._element.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))


def _set_cell_margins(cell, top=80, bottom=80, left=90, right=90) -> None:
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    )


def _add_text(paragraph, text: str, *, bold=False, size=9, color="334155"):
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    return run


def _add_table(doc, headers: list[str], rows: list[list[str]], right_columns=()) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for column, header in enumerate(headers):
        cell = table.cell(0, column)
        _set_cell_background(cell, "F1F5F9")
        _set_cell_margins(cell)
        _add_text(cell.paragraphs[0], header, bold=True, size=8.5)
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            if column in right_columns:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _add_text(paragraph, value, size=8.5)


def _format_quantity(value: Any) -> str:
    if value is None:
        return "—"
    try:
        return f"{float(value):,.0f} 台"
    except (TypeError, ValueError):
        return "—"


def _format_ratio(value: Any) -> str:
    if value is None:
        return "—"
    try:
        number = float(value) * 100
        return f"{'+' if number > 0 else ''}{number:.1f}%"
    except (TypeError, ValueError):
        return "—"


def _check_result(check: dict[str, Any]) -> str:
    if check.get("evaluable") is False:
        return "不可评估"
    return {"low": "正常", "medium": "关注", "high": "重点偏离"}.get(check.get("level"), "未分级")


def _finding_text(check: dict[str, Any]) -> str:
    findings = check.get("findings") or []
    messages = [str(item.get("message")) for item in findings if item.get("message")]
    return "；".join(messages) or "未提供命中说明"


def build_sop_report_docx(report_data: dict[str, Any]) -> io.BytesIO:
    """Build a Word report without inventing recommendations or business facts."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    spu = report_data.get("spu") or {}
    spu_code = spu.get("spu_code") or "—"
    spu_name = spu.get("spu_name") or spu_code
    as_of_date = str(report_data.get("as_of_date") or "—")
    readiness = get_report_readiness_label(report_data.get("completeness_status", "not_ready"))

    title = doc.add_paragraph()
    _add_text(title, f"{spu_name} S&OP 会议事实报告", bold=True, size=18, color="172033")
    subtitle = doc.add_paragraph()
    _add_text(
        subtitle,
        f"SPU：{spu_code} · 数据截至：{as_of_date} · 报告状态：{readiness}",
        size=9.5,
        color="64748B",
    )

    meta_rows = [
        ["所属产品线", spu.get("product_line") or "—", "生命周期", spu.get("lifecycle_stage") or "—"],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M"), "报告性质", "事实与规则校验结果"],
    ]
    _add_table(doc, ["字段", "内容", "字段", "内容"], meta_rows)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    _add_text(
        note,
        "说明：本报告仅汇总接口返回的业务事实和确定性规则校验结果，不自动生成供需情景、排产建议、责任人或行动状态。",
        size=9,
        color="475569",
    )

    doc.add_heading("一、预测校验明细", level=1)
    checks = report_data.get("forecast_checks") or []
    check_rows = []
    for check in checks:
        scope = f"{check.get('region') or '全部区域'} / {check.get('channel') or '全部渠道'}"
        check_rows.append(
            [
                str(check.get("forecast_month") or "—")[:7],
                scope,
                _format_quantity(check.get("forecast_qty")),
                _format_quantity(check.get("baseline_qty")),
                _format_ratio(check.get("deviation_ratio")),
                _check_result(check),
                str(check.get("rule_version") or "—"),
                _finding_text(check),
            ]
        )
    if check_rows:
        _add_table(
            doc,
            ["月份", "区域 / 渠道", "提报预测", "校验基线", "偏差率", "结果", "规则版本", "命中说明"],
            check_rows,
            right_columns=(2, 3, 4),
        )
    else:
        _add_text(doc.add_paragraph(), "暂无预测校验记录。", color="64748B")

    doc.add_heading("二、月度出库与终端激活", level=1)
    actuals = report_data.get("monthly_actuals") or []
    actual_rows = []
    for actual in actuals:
        outbound = actual.get("outbound_qty")
        activation = actual.get("activation_qty")
        diff = None if outbound is None or activation is None else float(outbound) - float(activation)
        actual_rows.append(
            [
                str(actual.get("period") or "—"),
                _format_quantity(outbound),
                _format_quantity(activation),
                _format_quantity(diff),
            ]
        )
    if actual_rows:
        _add_table(doc, ["月份", "出库量", "终端激活量", "数量差额"], actual_rows, right_columns=(1, 2, 3))
    else:
        _add_text(doc.add_paragraph(), "暂无月度出库与终端激活记录。", color="64748B")

    doc.add_heading("三、待补充决策数据", level=1)
    _add_text(
        doc.add_paragraph(),
        "供需情景与滚动计划暂不可评估。当前报告契约未提供库存、物料供应、产能负荷、计划版本、审批人和行动项数据。",
        color="64748B",
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
