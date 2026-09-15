import unittest

from docx import Document

from app.modules.sop.domain import docx_exporter


class SopDocxExporterTests(unittest.TestCase):
    def test_report_readiness_labels_distinguish_all_contract_states(self):
        get_label = getattr(docx_exporter, "get_report_readiness_label", lambda _status: None)

        self.assertEqual(get_label("complete"), "核心数据已就绪（可评审）")
        self.assertEqual(get_label("partial"), "数据域不完整（待补充）")
        self.assertEqual(get_label("not_ready"), "核心数据未就绪（暂不可评审）")

    def test_report_contains_only_contract_facts_and_explicit_unavailable_sections(self):
        report = {
            "spu": {"spu_code": "C706", "spu_name": "C706"},
            "as_of_date": "2026-09-30",
            "completeness_status": "partial",
            "forecast_checks": [
                {
                    "forecast_month": "2026-10-01",
                    "region": "华东",
                    "channel": "线上",
                    "forecast_qty": 120,
                    "baseline_qty": 100,
                    "deviation_ratio": 0.2,
                    "score": 80,
                    "evaluable": True,
                    "level": "low",
                    "rule_version": "forecast-check.v0.1-draft",
                    "findings": [{"message": "当前确定性规则未发现明显异常。"}],
                }
            ],
            "monthly_actuals": [{"period": "2026-08", "outbound_qty": 90, "activation_qty": 80}],
        }

        document = Document(docx_exporter.build_sop_report_docx(report))
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])

        self.assertIn("当前确定性规则未发现明显异常。", text)
        self.assertIn("forecast-check.v0.1-draft", text)
        self.assertIn("供需情景与滚动计划暂不可评估", text)
        for fabricated in ["42 天", "82%", "齐套率 98%", "CEO/VP", "悲观情景", "推荐基线"]:
            self.assertNotIn(fabricated, text)


if __name__ == "__main__":
    unittest.main()
